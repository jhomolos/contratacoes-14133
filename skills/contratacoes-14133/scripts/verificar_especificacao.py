#!/usr/bin/env python3
"""Verifica especificações técnicas (Lei nº 14.133/2021) quanto ao formato e a erros comuns.

Formato esperado:
  [Especificação]. Referência: [marca/modelo], similar ou superior.   -> modo "similar" (variante A)
  [Especificação]. Referência: [marca/modelo].                        -> modo "exigida" (variante B)
  [Especificação].                                                    -> modo "sem"     (variante C)

Uso:
  python verificar_especificacao.py --modo similar "texto"
  python verificar_especificacao.py --arquivo specs.txt [--modo auto]          (uma especificação por linha)
  python verificar_especificacao.py --xlsx arq.xlsx --coluna "ESPECIFICAÇÃO" [--aba NOME] [--modo-coluna "VARIANTE"]

Saída: ERRO (corrigir antes de publicar), AVISO (avaliar) e OK. O código de retorno é 1 se houver ERRO.
"""
import argparse, re, sys

REF_SIM = re.compile(r"Referência:\s*(?P<ref>[^\n]+?),\s*similar ou superior\.\s*$")
REF_EXI = re.compile(r"Referência:\s*(?P<ref>[^\n]+?)\.\s*$")

# (padrão, mensagem, severidade)
REGRAS = [
    (r"incerteza[^.;]{0,40}\bm[ií]nima\b", "Incerteza 'mínima': o sentido provavelmente está invertido (a incerteza é um limite máximo).", "ERRO"),
    (r"perda de retorno[^.;]{0,30}\bm[aá]xima\b", "Perda de retorno 'máxima': o sentido está invertido (deve ser mínima).", "ERRO"),
    (r"isola[çc][ãa]o[^.;]{0,30}\bm[aá]xima\b", "Isolação 'máxima': o sentido está invertido (deve ser mínima).", "ERRO"),
    (r"(tempo de subida|tempo de descida|rise ?time)[^.;]{0,30}(de no m[ií]nimo|m[ií]nimo de)", "Tempo de subida/descida 'no mínimo': o sentido está invertido (é um limite máximo).", "ERRO"),
    (r"(SWR|VSWR|ROE)[^.;]{0,25}\bm[ií]nim[oa]\b", "SWR 'mínimo': o sentido está invertido (é um limite máximo).", "ERRO"),
    (r"(ru[ií]do|ondula[çc][ãa]o|jitter)[^.;]{0,25}\bm[ií]nim[oa]\b", "Ruído/ondulação/jitter 'mínimo': o sentido está invertido.", "ERRO"),
    (r"\bt[ií]pic[oa]s?\b", "Valor 'típico': só é aceitável se o usuário concordar; requisitos metrológicos pedem valor garantido.", "AVISO"),
    (r"\bprecis[ãa]o\b", "'Precisão': prefira exatidão, incerteza ou resolução (VIM).", "AVISO"),
    (r"\b(chin[eê]s|chinesa|china|origem|nacionalidade|importad[oa]|fabrica[çc][ãa]o nacional|pa[ií]s de fabrica)", "Menção a origem ou nacionalidade: excluir por origem é vedado (art. 9º, I). Revise.", "ERRO"),
    (r"(\[DEFINIR|\[CONFIRMAR|\[VALIDAR|___|⚠)", "Há marcadores pendentes ([DEFINIR], [CONFIRMAR], [VALIDAR], ___ ou ⚠).", "ERRO"),
    (r"de (alta|boa|primeira) qualidade|robust[oa]|de primeira linha|excelente", "Adjetivo vago, não verificável.", "AVISO"),
    (r"ou equivalente", "'ou equivalente' junto com a variante de referência: use somente ', similar ou superior' no final.", "AVISO"),
]


def detectar_modo(texto):
    t = texto.strip()
    if REF_SIM.search(t):
        return "similar"
    if "Referência:" in t:
        return "exigida"
    return "sem"


def verificar(texto, modo="auto"):
    t = " ".join(str(texto).split())
    achados = []
    if not t:
        return [("ERRO", "Especificação vazia.")]
    if not t.endswith("."):
        achados.append(("ERRO", "O texto deve terminar com ponto final."))
    modo_det = detectar_modo(t)
    if modo == "auto":
        modo = modo_det
    n_ref = t.count("Referência:")
    if n_ref > 1:
        achados.append(("ERRO", "Há mais de um 'Referência:'; deve haver no máximo um, no final."))
    if modo == "similar":
        if not REF_SIM.search(t):
            achados.append(("ERRO", "Variante A: o texto deve terminar com 'Referência: <marca/modelo>, similar ou superior.'"))
    elif modo == "exigida":
        if REF_SIM.search(t) or "similar ou superior" in t:
            achados.append(("ERRO", "Variante B (marca exigida): não use ', similar ou superior'."))
        elif not REF_EXI.search(t):
            achados.append(("ERRO", "Variante B: o texto deve terminar com 'Referência: <marca/modelo>.'"))
    elif modo == "sem":
        if n_ref:
            achados.append(("ERRO", "Variante C (sem referência): o texto não deve conter 'Referência:'."))
    else:
        achados.append(("ERRO", f"Modo desconhecido: {modo}"))
    if n_ref == 1:
        cauda = t[t.find("Referência:"):]
        if re.search(r"\.\s+[A-ZÁÉÍÓÚÂÊÔÃÕÇ]", cauda):
            achados.append(("ERRO", "A referência deve ser a última frase do texto."))
    for padrao, msg, sev in REGRAS:
        if re.search(padrao, t, flags=re.IGNORECASE):
            achados.append((sev, msg))
    if len(t) < 60:
        achados.append(("AVISO", "Especificação muito curta; verifique se as características mínimas estão descritas."))
    return achados or [("OK", f"Sem problemas detectados (variante: {modo}).")]


def imprimir(rotulo, achados):
    print(f"--- {rotulo}")
    for sev, msg in achados:
        print(f"  [{sev}] {msg}")


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")  # evita acentos corrompidos no console do Windows
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("texto", nargs="?", help="Especificação a verificar")
    ap.add_argument("--modo", default="auto", choices=["auto", "similar", "exigida", "sem"])
    ap.add_argument("--arquivo", help="Arquivo .txt com uma especificação por linha")
    ap.add_argument("--xlsx", help="Planilha .xlsx")
    ap.add_argument("--aba", help="Nome da aba (padrão: a primeira)")
    ap.add_argument("--coluna", help="Cabeçalho da coluna com as especificações")
    ap.add_argument("--modo-coluna", help="Cabeçalho da coluna com a variante (similar/exigida/sem ou A/B/C)")
    a = ap.parse_args()

    itens = []  # (rótulo, texto, modo)
    if a.texto:
        itens.append(("texto", a.texto, a.modo))
    if a.arquivo:
        with open(a.arquivo, encoding="utf-8") as f:
            for i, linha in enumerate(f, 1):
                if linha.strip():
                    itens.append((f"linha {i}", linha, a.modo))
    if a.xlsx:
        import openpyxl
        wb = openpyxl.load_workbook(a.xlsx, data_only=True)
        ws = wb[a.aba] if a.aba else wb.worksheets[0]
        cab = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        if a.coluna not in cab:
            sys.exit(f"Coluna '{a.coluna}' não encontrada. Cabeçalhos: {cab}")
        ci = cab.index(a.coluna)
        mi = cab.index(a.modo_coluna) if a.modo_coluna and a.modo_coluna in cab else None
        mapa = {"A": "similar", "B": "exigida", "C": "sem"}
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[ci]:
                modo = a.modo
                if mi is not None and r[mi]:
                    v = str(r[mi]).strip()
                    modo = mapa.get(v.upper(), v.lower())
                itens.append((f"linha {len(itens) + 2}", r[ci], modo))
    if not itens:
        ap.print_help()
        sys.exit(2)

    erro = False
    for rot, txt, modo in itens:
        ach = verificar(txt, modo)
        erro |= any(s == "ERRO" for s, _ in ach)
        imprimir(rot, ach)
    sys.exit(1 if erro else 0)


if __name__ == "__main__":
    main()
